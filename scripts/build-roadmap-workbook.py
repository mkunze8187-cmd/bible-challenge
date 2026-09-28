"""Build outputs/roadmap-visual/bible-challenge-roadmap.xlsx from scripts/roadmap-data.json.

Python/openpyxl port of build-roadmap-workbook.mjs (same sheets, layout and colors) for
environments without @oai/artifact-tool. Optional: --previews renders PNG previews via
Microsoft Excel (COM) + PyMuPDF on Windows.

Usage (from repo root): python scripts/build-roadmap-workbook.py [--previews]
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / "scripts" / "roadmap-data.json").read_text(encoding="utf-8"))
OUT_DIR = ROOT / "outputs" / "roadmap-visual"
OUT = OUT_DIR / "bible-challenge-roadmap.xlsx"

FONT = "Arial"
C = {
    "navy": "1E3A5F", "blue": "2F6F9F", "teal": "2A7F7F", "green": "4E8B5A", "amber": "B7791F",
    "red": "B84A4A", "purple": "7251A5", "gray": "5F6B7A", "lightGray": "EEF1F4",
    "header": "233142", "text": "202B36", "border": "C9D2DC",
}
THIN = Side(style="thin", color=C["border"])

phases = DATA["phases"]
milestones = DATA["milestones"]
issue_rows = DATA["issueRows"]
step_labels = DATA["stepLabels"]
FIRST_STEP_COL = 6  # F (1-based)
LAST_STEP_COL = FIRST_STEP_COL + len(step_labels) - 1
NOTE_COL = LAST_STEP_COL + 1
L = get_column_letter


def base_font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), color=kw.pop("color", C["text"]), **kw)


def setup(ws):
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=1, max_row=60, max_col=20):
        for cell in row:
            cell.font = base_font()
            cell.alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 2


def title(ws, text, subtitle=""):
    ws["B2"] = text
    ws["B2"].font = base_font(size=16, bold=True)
    ws["B3"] = subtitle
    ws["B3"].font = base_font(italic=True, color=C["gray"])
    for col in range(2, 16):
        ws.cell(row=4, column=col).border = Border(bottom=THIN)


def write(ws, top, left, rows):
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            ws.cell(row=top + r, column=left + c, value=v)


def header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=C["header"])
        cell.font = base_font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def outline(ws, r1, c1, r2, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = Border(
                left=THIN if c == c1 else cell.border.left,
                right=THIN if c == c2 else cell.border.right,
                top=THIN if r == r1 else cell.border.top,
                bottom=THIN if r == r2 else cell.border.bottom,
            )


def wrap(ws, r1, c1, r2, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal=cell.alignment.horizontal)


def bar_chart(title_text, color, data_ref, cats_ref, width, height):
    ch = BarChart()
    ch.type = "col"
    ch.title = title_text
    ch.legend = None
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    ch.series[0].graphicalProperties.solidFill = color
    ch.series[0].graphicalProperties.line.solidFill = color
    ch.y_axis.numFmt = "0"
    ch.y_axis.majorGridlines = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.width, ch.height = width, height
    return ch


wb = Workbook()
dash = wb.active
dash.title = "Dashboard"
ms = wb.create_sheet("Milestones")
im = wb.create_sheet("Issue Map")
cd = wb.create_sheet("Chart Data")
for ws, tab in ((dash, "navy"), (ms, "blue"), (im, "teal"), (cd, "lightGray")):
    ws.sheet_properties.tabColor = C[tab]
    setup(ws)

# Dashboard
title(dash, DATA["title"], DATA["subtitle"])
write(dash, 6, 2, DATA["nearTerm"])
header(dash, 6, 2, 5)
outline(dash, 6, 2, 8, 5)
write(dash, 6, 7, DATA["releaseGate"])
wrap(dash, 6, 2, 8, 5)
wrap(dash, 6, 7, 8, 10)
for r in (6, 7, 8):
    dash.row_dimensions[r].height = 44
header(dash, 6, 7, 10)
outline(dash, 6, 7, 8, 10)

write(dash, 11, 2, [["Track", "Milestone", "Issues", "Status", *step_labels, "Dependency note"]])
header(dash, 11, 2, NOTE_COL)
for i, p in enumerate(phases):
    row = 12 + i
    marks = ["■" if p["start"] <= s + 1 < p["start"] + p["duration"] else "" for s in range(len(step_labels))]
    write(dash, row, 2, [[p["title"], p["milestone"], p["issues"], p["status"], *marks, p["dependency"]]])
    color = C[p["color"]]
    dash.cell(row=row, column=2).font = base_font(bold=True)
    dash.cell(row=row, column=4).alignment = Alignment(vertical="center", wrap_text=True)
    for c in range(5, LAST_STEP_COL + 1):
        cell = dash.cell(row=row, column=c)
        cell.font = base_font(size=13 if c > 5 else 11, bold=True, color=color)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dash.cell(row=row, column=NOTE_COL).alignment = Alignment(vertical="center", wrap_text=True)
    dash.row_dimensions[row].height = 42
last_phase_row = 11 + len(phases)
outline(dash, 11, 2, last_phase_row, NOTE_COL)

note_row = last_phase_row + 3
dash.cell(row=note_row, column=2, value=DATA["readingNote"])
dash.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=NOTE_COL)
dash.cell(row=note_row, column=2).font = base_font(italic=True, color=C["gray"])
dash.cell(row=note_row, column=2).alignment = Alignment(vertical="center", wrap_text=True)
dash.row_dimensions[note_row].height = 42

chart_phases = [p for p in phases if p["open"] > 0]
status_row = note_row + 3
write(dash, status_row, 2, [["Track", "Open issues"], *[[p["title"], p["open"]] for p in chart_phases]])
header(dash, status_row, 2, 3)
outline(dash, status_row, 2, status_row + len(chart_phases), 3)
dash.add_chart(
    bar_chart(
        "Open issues by roadmap track", C["blue"],
        Reference(dash, min_col=3, min_row=status_row, max_row=status_row + len(chart_phases)),
        Reference(dash, min_col=2, min_row=status_row + 1, max_row=status_row + len(chart_phases)),
        30, 11,
    ),
    f"E{status_row}",
)

widths = {"B": 35, "C": 18, "D": 28, "E": 24}
for col, w in widths.items():
    dash.column_dimensions[col].width = w
for c in range(FIRST_STEP_COL, LAST_STEP_COL + 1):
    dash.column_dimensions[L(c)].width = 10
dash.column_dimensions[L(NOTE_COL)].width = 58

# Milestones
title(ms, "Milestone Summary", "Open/closed counts and planning notes by GitHub milestone.")
write(ms, 6, 2, [["Milestone", "Issues", "Open", "Closed", "Planning note"]])
header(ms, 6, 2, 6)
write(ms, 7, 2, milestones)
last_ms = 6 + len(milestones)
outline(ms, 6, 2, last_ms, 6)
wrap(ms, 7, 3, last_ms, 3)
wrap(ms, 7, 6, last_ms, 6)
ms.column_dimensions["B"].width = 46
ms.column_dimensions["C"].width = 30
ms.column_dimensions["D"].width = 8
ms.column_dimensions["E"].width = 8
ms.column_dimensions["F"].width = 58
ms.freeze_panes = "A7"
chart_top = last_ms + 3
write(ms, chart_top, 2, [["Milestone", "Open"], *[[m[0], m[2]] for m in milestones]])
ms.add_chart(
    bar_chart(
        "Open issues by milestone", C["teal"],
        Reference(ms, min_col=3, min_row=chart_top, max_row=chart_top + len(milestones)),
        Reference(ms, min_col=2, min_row=chart_top + 1, max_row=chart_top + len(milestones)),
        30, 13,
    ),
    "H6",
)

# Issue Map
title(im, "Issue Relationship Map", "Grouped issue references, blockers, and planning action.")
write(im, 6, 2, [["Issues", "Work item", "Milestone", "Status", "Depends on", "Planning action"]])
header(im, 6, 2, 7)
write(im, 7, 2, issue_rows)
last_im = 6 + len(issue_rows)
outline(im, 6, 2, last_im, 7)
wrap(im, 7, 2, last_im, 7)
for col, w in {"B": 26, "C": 44, "D": 12, "E": 10, "F": 34, "G": 26}.items():
    im.column_dimensions[col].width = w
im.freeze_panes = "A7"

# Chart Data
title(cd, "Chart Data", "Source data used by the roadmap charts.")
write(cd, 6, 2, [["Order", "Milestone", "Track", "Start step", "Duration", "Open", "Closed", "Issues", "Theme", "Dependency note"]])
header(cd, 6, 2, 11)
write(cd, 7, 2, [[p["order"], p["milestone"], p["title"], p["start"], p["duration"], p["open"], p["closed"], p["issues"], p["theme"], p["dependency"]] for p in phases])
last_cd = 6 + len(phases)
outline(cd, 6, 2, last_cd, 11)
wrap(cd, 7, 9, last_cd, 11)
for col, w in {"B": 8, "C": 13, "D": 42, "E": 10, "F": 10, "G": 8, "H": 8, "I": 34, "J": 16, "K": 65}.items():
    cd.column_dimensions[col].width = w
cd.freeze_panes = "A7"

for ws in (dash, ms, im, cd):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

OUT_DIR.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"saved={OUT}")

# Inspection dump (tables as values) for diff-friendly review.
lines = [json.dumps({"kind": "workbook", "sheets": len(wb.worksheets), "generated": DATA["generated"]})]
for ws in wb.worksheets:
    values = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)]
    lines.append(json.dumps({"kind": "table", "sheet": ws.title, "address": f"A1:{L(ws.max_column)}{ws.max_row}", "values": values}, ensure_ascii=False))
(OUT_DIR / "bible-challenge-roadmap.xlsx.inspect.ndjson").write_text("\n".join(lines) + "\n", encoding="utf-8")

if "--previews" in sys.argv:
    import subprocess
    import tempfile

    import pymupdf as fitz

    previews = (("Dashboard", "dashboard-preview.png"), ("Milestones", "milestones-preview.png"), ("Issue Map", "issue-map-preview.png"))
    tmp = Path(tempfile.gettempdir())
    exports = "; ".join(
        f"$b.Worksheets.Item('{name}').ExportAsFixedFormat(0, '{tmp / ('roadmap-' + name.replace(' ', '-') + '.pdf')}')" for name, _ in previews
    )
    ps = (
        "$x = New-Object -ComObject Excel.Application; $x.Visible = $false; $x.DisplayAlerts = $false; "
        f"try {{ $b = $x.Workbooks.Open('{OUT}'); {exports}; $b.Close($false) }} finally {{ $x.Quit() }}"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    for name, png in previews:
        pdf = tmp / f"roadmap-{name.replace(' ', '-')}.pdf"
        doc = fitz.open(pdf)
        page = doc[0]
        clip = fitz.Rect()
        for _, rect in page.get_bboxlog():
            clip |= fitz.Rect(rect)
        page.get_pixmap(dpi=200, clip=clip + (-12, -12, 12, 12)).save(OUT_DIR / png)
        doc.close()
        pdf.unlink()
    print("previews rendered")
